from pathlib import Path
from datetime import datetime, timezone
import json, math, textwrap
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / 'assets'
ASSETS.mkdir(parents=True, exist_ok=True)
BASE = ROOT.parents[1]
TDD = json.loads((BASE / 'blueprint' / 'technical-design.json').read_text(encoding='utf-8'))
INK = '#122A3A'
TEAL = '#087E82'
PALE = '#EAF5F4'
LINE = '#C8D9DF'
GRAY = '#526775'
FONT_DIR = Path('C:/Windows/Fonts')

def font(size, bold=False):
    return ImageFont.truetype(str(FONT_DIR / ('seguisb.ttf' if bold else 'segoeui.ttf')), size)

def canvas(name, width=1500, height=760):
    im = Image.new('RGB', (width, height), 'white')
    return im, ImageDraw.Draw(im)

def txt(draw, xy, value, size=26, fill=INK, bold=False):
    draw.text(xy, value, font=font(size, bold), fill=fill)

def wrap(draw, value, width, size=25, bold=False):
    words = value.split()
    lines, current = [], ''
    for word in words:
        candidate = (current + ' ' + word).strip()
        if draw.textlength(candidate, font=font(size, bold)) > width and current:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines

def block(draw, x, y, w, h, title, body='', fill=PALE, label=None, size=27, body_size=23, gap=12):
    draw.rounded_rectangle((x,y,x+w,y+h), radius=18, fill=fill, outline=LINE, width=2)
    ty = y + 22
    if label:
        txt(draw, (x+24,ty), label, 18, TEAL, True)
        ty += 32
    for line in wrap(draw, title, w-48, size, True):
        txt(draw, (x+24,ty), line, size, INK, True)
        ty += size+9
    if body:
        ty += gap
        for line in wrap(draw, body, w-48, body_size):
            txt(draw, (x+24,ty), line, body_size, GRAY)
            ty += body_size+8

def arrow(draw, start, end, color=TEAL, width=5):
    draw.line((start,end), fill=color, width=width)
    angle = math.atan2(end[1]-start[1], end[0]-start[0])
    points = [end]
    for delta in (2.65,-2.65):
        points.append((end[0]+18*math.cos(angle+delta),end[1]+18*math.sin(angle+delta)))
    draw.polygon(points, fill=color)

def save(im, name):
    path = ASSETS / (name + '.png')
    im.save(path)
    return path

def make_visuals():
    im,d = canvas('cover',1500,700)
    d.rounded_rectangle((25,30,1080,620),24,fill='#F2F6F8',outline=LINE,width=3)
    d.rounded_rectangle((50,60,1055,130),14,fill=INK)
    txt(d,(78,80),'Agent8lite',30,'white',True)
    txt(d,(810,84),'WORKSPACE',20,'white',True)
    d.rectangle((50,150,235,595),fill='white')
    for j,s in enumerate(['Overview','Contacts','Properties','Website','Marketing','Reports']):
        txt(d,(76,180+j*54),s,23,TEAL if j==0 else GRAY,j==0)
    txt(d,(265,168),'Your business at a glance',33,INK,True)
    for j,(t,b) in enumerate([('Follow up','Assigned leads'),('Next actions','Tasks and bookings'),('Performance','Defined metrics')]):
        block(d,265+j*253,230,235,130,t,b,fill='white',size=25)
    block(d,265,390,735,172,'Lead activity','Capture an inquiry, assign an agent and track the next action.',fill='white',size=27)
    d.rounded_rectangle((1130,90,1455,650),40,fill=INK)
    d.rounded_rectangle((1143,107,1442,631),30,fill='white')
    d.rounded_rectangle((1225,117,1360,140),10,fill=INK)
    txt(d,(1168,173),'Agent8lite',27,INK,True)
    txt(d,(1168,222),'Today',35,INK,True)
    for j,(t,b) in enumerate([('Lead inquiry','Review and assign'),('Follow up task','Open contact record'),('Website update','Preview a draft')]):
        block(d,1164,284+j*100,254,90,t,b,fill=PALE,size=21,body_size=19,gap=6)
    save(im,'cover-experience')

    im,d=canvas('platform',1500,690)
    block(d,410,20,680,125,'Agent8lite business workspace','Shared identity, permissions, brand and activity',fill=INK,size=32)
    # Keep the central dark card readable with white text.
    d.rounded_rectangle((410,20,1090,145),18,fill=INK)
    txt(d,(450,43),'Agent8lite business workspace',32,'white',True)
    txt(d,(450,92),'Identity  /  Permissions  /  Brand  /  Activity',23,'white')
    groups=[('Attract','Websites, property search and business presence'),('Manage','Contacts, leads, tasks, deals and appointments'),('Engage','Campaigns, automation and reviewed AI drafts'),('Measure','Reporting, account health and operational oversight')]
    for i,(t,b) in enumerate(groups):
        x=20+i*375
        arrow(d,(750,145),(x+175,235))
        block(d,x,250,350,230,t,b,fill=PALE,size=31)
    d.rounded_rectangle((20,535,1475,648),18,fill='#F2F6F8',outline=LINE,width=2)
    txt(d,(48,558),'Desktop browser  /  iPhone home screen  /  Android home screen',29,INK,True)
    txt(d,(48,605),'One product experience with secure access and complete core workflows',24,GRAY)
    save(im,'platform-map')

    im,d=canvas('journey',1500,510)
    stages=[('Discover','Website and property search'),('Inquire','Form and contact preference'),('Assign','Eligible agent and routing rule'),('Follow up','Task, booking or campaign'),('Measure','Pipeline and source reporting')]
    for i,(t,b) in enumerate(stages):
        x=15+i*300
        block(d,x,80,270,225,t,b,label=f'0{i+1}',size=30)
        if i<4: arrow(d,(x+271,192),(x+293,192))
    d.rounded_rectangle((15,355,1485,475),18,fill='#F2F6F8',outline=LINE,width=2)
    txt(d,(44,380),'A shared record preserves the source, responsible agent and next action.',28,INK,True)
    txt(d,(44,425),'Every handoff respects organization access, consent and permitted listing use.',24,GRAY)
    save(im,'customer-journey')

    im,d=canvas('crm',1500,370)
    for i,(t,b) in enumerate([('Contact','Person and channels'),('Lead','Intent and follow up'),('Deal','Pipeline opportunity')]):
        x=40+i*500
        block(d,x,35,420,205,t,b,fill=PALE,size=34)
        if i<2: arrow(d,(x+430,137),(x+485,137))
    txt(d,(40,288),'Activity, consent, ownership and stage history remain connected.',27,GRAY)
    save(im,'crm-records')

    im,d=canvas('mls',1500,520)
    stages=[('Trestle feed','Approved source access'),('Data policy','Fields, media and rights'),('Property experience','Search, detail and organizer'),('Agent workflow','Inquiry, routing and follow up')]
    for i,(t,b) in enumerate(stages):
        x=20+i*375
        block(d,x,60,340,240,t,b,label=['SOURCE','CONTROL','CUSTOMER','BUSINESS'][i],size=29)
        if i<3: arrow(d,(x+341,182),(x+369,182))
    block(d,20,360,1460,120,'Live activation requires an approved Las Vegas feed','Display, images, freshness, alerts, retention and AI use are reviewed independently.',fill='#F2F6F8',size=26)
    save(im,'mls-experience')

    im,d=canvas('marketing',1500,400)
    stages=[('Design','Message and audience'),('Review','Preview and confirm'),('Check','Consent and readiness'),('Execute','Schedule and delivery'),('Reconcile','Receipts and outcomes')]
    for i,(t,b) in enumerate(stages):
        x=15+i*300
        block(d,x,55,270,205,t,b,label=f'0{i+1}',size=29)
        if i<4: arrow(d,(x+271,154),(x+293,154))
    txt(d,(25,318),'Eligibility is checked again when a scheduled action is due.',28,GRAY)
    save(im,'marketing-flow')

    im,d=canvas('ai',1500,365)
    for i,(t,b) in enumerate([('Use permitted facts','Scope and source checks'),('Generate a proposal','Draft, summary or explanation'),('Human review','Edit, accept or reject'),('Save a draft','Publish or send separately')]):
        x=20+i*375
        block(d,x,35,340,210,t,b,fill=PALE,size=28)
        if i<3: arrow(d,(x+341,144),(x+369,144))
    txt(d,(25,295),'People retain control of publication, communication and commercial actions.',27,GRAY)
    save(im,'ai-review')

    im,d=canvas('mobile',1500,720)
    for i,(t,items) in enumerate([('Contacts',['Search contacts','View activity','Edit a lead','Complete a task']),('Properties',['Filter results','Open listing detail','Save a property','Submit an inquiry']),('Business',['Preview a page','Review a campaign','Read a report','Manage settings'])]):
        x=120+i*455
        d.rounded_rectangle((x,20,x+350,650),42,fill=INK)
        d.rounded_rectangle((x+14,35,x+336,630),31,fill='white')
        d.rounded_rectangle((x+111,44,x+239,69),12,fill=INK)
        txt(d,(x+36,104),'Agent8lite',27,INK,True)
        txt(d,(x+36,163),t,34,INK,True)
        for j,item in enumerate(items):
            d.rounded_rectangle((x+34,242+j*77,x+316,302+j*77),12,fill=PALE,outline=LINE,width=2)
            txt(d,(x+50,256+j*77),item,27,INK)
        d.line((x+106,605,x+244,605),fill=INK,width=6)
    txt(d,(120,678),'Browser and installed modes use the same permissions and business records.',25,GRAY)
    save(im,'mobile-experience')

    im,d=canvas('roadmap',1500,850)
    phases=[('M1','Foundation','Secure accounts and shared standards'),('M2','Core workspace','Accounts, CRM and agent activity'),('M3','Public websites','Content, domains and lead capture'),('M4','Marketing','Campaigns and controlled automation'),('M5','Property preparation','Synthetic search and inquiry workflows'),('M6','Live Trestle','Licensed feed and partner acceptance'),('M7','Intelligence','Reviewed AI and advanced reporting'),('M8','Controlled rollout','Device, recovery and operational evidence')]
    for i,(ident,title,body) in enumerate(phases):
        row=i//4; col=i%4; x=20+col*375; y=25+row*375
        block(d,x,y,340,285,title,body,label=ident,size=29)
        if col<3: arrow(d,(x+341,y+147),(x+369,y+147))
    txt(d,(25,780),'Milestones define acceptance order. Independent work can proceed in parallel.',26,GRAY)
    save(im,'delivery-roadmap')

make_visuals()
doc = Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.62)
sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
sec.header_distance=Inches(.25); sec.footer_distance=Inches(.28)
sec.different_first_page_header_footer=True
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    s=doc.styles[name]
    s.font.name='Calibri'
    s.font.italic=False
    s.font.color.rgb=RGBColor(0,0,0)
    s._element.rPr.rFonts.set(qn('w:ascii'),'Calibri')
    s._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri')
    s.paragraph_format.space_after=Pt(7)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.10
doc.styles['Title'].font.size=Pt(32)
doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(13)
doc.styles['Subtitle'].font.size=Pt(17)
doc.styles['Subtitle'].font.bold=False
doc.styles['Heading 1'].font.size=Pt(22)
doc.styles['Heading 1'].font.bold=True
doc.styles['Heading 1'].paragraph_format.space_after=Pt(12)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].font.bold=True
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(5)
doc.styles['Caption'].font.size=Pt(8.5)
doc.styles['Caption'].font.bold=False
doc.styles['Caption'].font.color.rgb=RGBColor.from_string('526775')
doc.styles['Caption'].paragraph_format.space_after=Pt(9)
header=sec.header.paragraphs[0]
header.text='Agent8lite   |   Application requirements and features'
header.style=doc.styles['Caption']
header.runs[0].font.color.rgb=RGBColor(0,0,0)
for footer in [sec.footer,sec.first_page_footer]:
    p=footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=p.add_run('Agent8lite  |  Product requirements 1.0  |  ');r.font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
    p._p.append(field)
doc.core_properties.title='Agent8lite Application Requirements and Feature Overview'
doc.core_properties.subject='Executive product requirements and feature specification'
doc.core_properties.author='Agent8lite'
doc.core_properties.keywords='Real estate, SaaS, application requirements, CRM, Trestle, PWA'
doc.core_properties.comments=''
doc.core_properties.created=datetime(2026,10,10,tzinfo=timezone.utc)
doc.core_properties.modified=datetime(2026,10,10,tzinfo=timezone.utc)

page_titles=[]
req_ids=[]
def page(title, kicker=None):
    if page_titles: doc.add_page_break()
    page_titles.append(title)
    if kicker:
        p=doc.add_paragraph(kicker.upper());p.paragraph_format.space_after=Pt(7)
        for r in p.runs:r.font.size=Pt(9);r.font.bold=True
    doc.add_paragraph(title,'Heading 1')

def p(text, bold_start=None):
    para=doc.add_paragraph()
    if bold_start:
        para.add_run(bold_start+' ').bold=True
    para.add_run(text)
    return para

def h(text): doc.add_paragraph(text,'Heading 2')

def bullet(text):
    para=doc.add_paragraph(text,'List Bullet')
    para.paragraph_format.space_after=Pt(5)
    return para

def req(ident,title,body):
    req_ids.append(ident)
    para=doc.add_paragraph()
    para.paragraph_format.space_after=Pt(3 if page_titles[-1].startswith('A complete application') else 6)
    para.add_run(title+'. ').bold=True
    para.add_run(body+' ')
    r=para.add_run('['+ident+']');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('526775')
    return para

def figure(name, caption, width=7.06):
    para=doc.add_paragraph()
    para.paragraph_format.space_after=Pt(4)
    para.paragraph_format.keep_with_next=True
    shape=para.add_run().add_picture(str(ASSETS/(name+'.png')),width=Inches(width))
    shape._inline.docPr.set('descr',caption)
    doc.add_paragraph(caption,'Caption')

def table(headers, rows, widths, compact=False):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    props.append(borders)
    margin=OxmlElement('w:tblCellMar')
    for edge in ['top','bottom','left','right']:
        e=OxmlElement('w:'+edge);e.set(qn('w:w'),('60' if compact else '100') if edge in ['top','bottom'] else '115');e.set(qn('w:type'),'dxa');margin.append(e)
    props.append(margin)
    for i,width in enumerate(widths):t.columns[i].width=Inches(width)
    for i,head in enumerate(headers):t.rows[0].cells[i].text=head
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for i,value in enumerate(row):cells[i].text=str(value)
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'122A3A' if ri==0 else ('F3F7F8' if ri%2==0 else 'FFFFFF'));cell._tc.get_or_add_tcPr().append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1 if compact else 2);para.paragraph_format.line_spacing=1.02 if compact else 1.05
                for r in para.runs:
                    r.font.size=Pt(8.7 if compact else 9)
                    r.font.bold=ri==0
                    r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
    spacer=doc.add_paragraph()
    spacer.paragraph_format.space_after=Pt(0)
    spacer.paragraph_format.line_spacing=Pt(3)
    spacer.add_run().font.size=Pt(3)
    return t

page_titles.append('Agent8lite Application Requirements and Feature Overview')
doc.add_paragraph('AGENT8LITE','Subtitle')
doc.add_paragraph('Application Requirements\nand Feature Overview','Title')
p('Executive product specification for a complete real estate business application.')
p('10 October 2026  |  Version 1.0')
figure('cover-experience','Conceptual product experience. Interface content is illustrative and does not represent an accepted production release.')
h('Purpose of this document')
p('Define the product capabilities, customer experience and launch conditions for Agent8lite. The platform brings websites, property discovery, lead management, marketing and business reporting into a shared workspace for agents and brokerages.')
p('The application must work in desktop browsers and as an installable home screen application on iPhone and Android, without app store distribution. Trestle is the selected MLS data provider, with live activation subject to approved feed access and usage rights.')

page('Executive brief','01 Executive overview')
p('Agent8lite is intended to give real estate professionals one place to attract inquiries, manage relationships, coordinate follow up and understand business performance. A shared record connects public website activity with the responsible agent, the next action and the resulting pipeline opportunity.')
h('Business outcomes')
bullet('Improve continuity between property discovery, inquiry capture, lead assignment and agent follow up.')
bullet('Reduce repeated data entry through connected contact, website, campaign and reporting workflows.')
bullet('Give brokerage leaders clear ownership, accountable actions and consistent reporting across teams.')
bullet('Support work in the field with complete core workflows on iPhone and Android.')
bullet('Establish secure organizational boundaries and controlled use of customer and licensed property data.')
h('Scope and product stage')
p('The defined product covers 11 capability areas and 76 functional requirements. Implementation has started. No delivery milestone is recorded as complete in the current project status. The requirements below describe the target product; release claims depend on the launch acceptance process.')
p('Commercial pricing, contractual service levels, launch dates and production operating costs remain business decisions. This document does not assign revenue forecasts or claim measured productivity improvements.')
h('Reading guide')
table(['Topic','Pages'],[
('Platform and customer journeys','3 to 4'),
('Functional requirements across all 11 areas','5 to 14'),
('Shared standards and launch scope','15 to 16'),
('Delivery roadmap and launch acceptance','17 to 18'),
('Business measures and executive decisions','19'),
('Requirement coverage and document governance','20')],[5.65,1.35])

page('The platform and its business role','02 Product model')
p('Public experiences create demand. The agent workspace manages the resulting relationships. Controlled communication supports follow up, while reporting makes progress and service health visible. Shared identity, permissions and brand settings connect these experiences across the product.')
figure('platform-map','Figure 1. Product capability map showing the four connected business activities and the common device experience.')
h('What the application must connect')
bullet('A website or property inquiry must become an accountable CRM record with source information and an eligible owner.')
bullet('Tasks, appointments and campaigns must use the same current contact, ownership and permission information.')
bullet('Leaders must be able to trace a reported result to its definition, time period and underlying activity.')
bullet('Changes in access, consent or listing rights must affect related screens, background work and downloads.')
p('The platform supports individual agents, offices and brokerages through organization and role controls. Feature access follows the customer subscription and the permissions granted to each user.')

page('Users and connected customer journeys','03 Users and experience')
p('The product must support public visitors, real estate professionals, organization administrators and authorized support staff. Each role sees only the functions and business records appropriate to its responsibilities.')
table(['User','Primary need','Access boundary'],[
('Public visitor','Find properties, inquire, save searches and manage preferences','Published content and permitted listings; own organizer'),
('Agent','Manage assigned leads, tasks, appointments and deals','Own or explicitly permitted team records'),
('Manager or owner','Allocate work, manage members and oversee the business','Granted team or organization scope'),
('Marketing editor','Manage content, brand assets and campaigns','Granted editing and publication rights'),
('Billing administrator','Review invoices and approve subscription changes','Billing permissions'),
('Read only analyst','Review business reports','Permitted report and source data scope'),
('Platform support','Resolve service and integration issues','Time limited, justified customer access')],[1.25,3.0,2.75])
figure('customer-journey','Figure 2. Inquiry journey from discovery to accountable follow up and reporting.')
p('Other essential journeys include onboarding and inviting a colleague, publishing a website revision, reviewing a scheduled campaign, resuming work on a phone and investigating a failed integration. Confirmation and recovery states must remain clear throughout each journey.')

page('Accounts and the agent workspace','04 Account and daily work requirements')
p('Agents need an actionable view of their work. Owners need predictable onboarding, permission management and subscription controls. These capabilities establish the shared context for every other feature.')
h('Organization and commercial management')
req('MOD-A-001','Organization setup','Create the organization, standard roles, initial pipeline and business profile. Show setup progress and failures. Invitations must expire, be single use and match the verified recipient.')
req('MOD-A-002','Roles and ownership','Provide six role presets and configurable access for own, team or organization records. Prevent removal of the final active owner and limit delegated permissions.')
req('MOD-A-003','Organization management','Manage offices, members and business profiles. Switching organizations must change the complete working context and respect visible subscription limits.')
req('MOD-A-004','Subscription experience','Show plans, enabled capabilities and invoices, distinguishing annual totals from monthly equivalents. Preview amounts, taxes, currency, timing and proration before a user confirms a billing change.')
req('MOD-A-005','Billing accuracy','Reconcile subscription status with the payment provider. A successful payment must not imply that unfinished account setup is complete.')
req('MOD-A-006','Platform administration','Provide account health and redacted audit views. Access to customer records requires a justified, time limited grant and a visible support session indicator.')
h('Daily agent workspace')
req('MOD-B-001','Prioritized work','Show assigned lead activity, overdue and upcoming tasks, appointments, pipeline, notifications and website or campaign health. Team views require explicit access.')
req('MOD-B-002','Trustworthy metrics','Show each metric definition, period, timezone, source and refresh time. Distinguish measured zero, loading, empty and unavailable states.')
req('MOD-B-003','Quick actions','Open existing records through shareable routes and offer authorized quick actions. Layout preferences must not change access rights.')
req('MOD-B-004','Appointments','Reflect verified booking and cancellation events. Display user and organization timezones clearly when they differ.')
req('MOD-B-005','Notifications','Provide a personal inbox. Reading notices must not change the related lead or task.')
h('Acceptance expectation')
p('An invited member must reach the correct workspace, organization switching must expose only permitted records, and dashboard totals must reconcile with the corresponding scoped report.')

page('CRM and accountable follow up','05 Relationship management requirements')
p('The CRM must preserve a clear relationship between the person, the inquiry and the business opportunity. Ownership, consent and the next action must remain visible as a lead moves through the pipeline.')
figure('crm-records','Figure 3. Contact, lead and deal records serve different business purposes while sharing relevant history.')
req('MOD-D-001','Customer record','Manage contact channels, source, voluntarily supplied intent and timeline, tags, notes, stage history and a communication timeline. Preserve unknown values accurately.')
req('MOD-D-002','Lists and pipeline','Provide saved filters, pagination, preferred columns, scoped bulk actions and a pipeline view. Start with Uncontacted, Attempted Contact, Nurturing, Dead, Under Contract and Closed; allow configurable stages.')
req('MOD-D-003','Lead routing','Apply eligible ownership and assignment rules, including first touch, listing agent, geography and weighted distribution. Explain the winning rule and retain an unassigned queue when necessary.')
req('MOD-D-004','Follow up prioritization','Use documented engagement signals and explain the score. Prioritization must not determine access to housing or financial eligibility.')
req('MOD-D-005','Reviewed imports','Provide field mapping, preview, row errors and duplicate candidates. Require a reviewed merge or skip decision before committing an import.')
req('MOD-D-006','Controlled exports and archive','Authorize exports at request and download, with expiring access and an audit record. Archiving must cancel eligible pending marketing and apply the retention policy.')
req('MOD-D-007','Tasks and history','Require an assignee, title and due time with timezone. Preserve change history and resolve competing edits without silently overwriting work.')
h('Acceptance expectation')
p('A repeated inquiry must create one accountable lead decision. Routing must exclude inactive or unauthorized agents. Import counts, consent history, task changes and exports must be explainable and consistent.')

page('Websites content and lead capture','06 Digital experience requirements')
p('Customers must be able to create an original branded real estate website and maintain it through a structured editor. The publishing process must keep drafts separate from public content and connect inquiries to the native CRM.')
req('MOD-C-001','Template families','Provide independently designed agent, brokerage and landing page templates. Preserve versioned templates and customer content separately.')
req('MOD-C-002','Content and page blocks','Support pages, blogs, landing pages, communities, agents and testimonials. Include navigation, SEO settings, social images and blocks for text, images, calls to action, forms and permitted property experiences.')
req('MOD-C-003','Draft preview publish and rollback','Save revisions, provide protected previews, validate content before publication and allow rollback to a compatible prior version. Competing edits must not lose content.')
req('MOD-C-004','Safe editing','Preserve unsupported content for recovery and explain migration needs. Initial editing must use registered blocks rather than unrestricted executable scripts or custom code.')
req('MOD-C-005','Domains and secure delivery','Verify domain ownership, assign each hostname uniquely and provision secure delivery. Domain removal must revoke the former site route before reassignment.')
req('MOD-C-006','Media management','Validate and scan uploads, create responsive images and record usage rights. Customer media and unpublished originals must remain appropriately protected.')
req('MOD-C-007','Forms and inquiry capture','Provide labeled fields, validation, abuse controls, purpose specific consent and confirmation. A submission must be durably recorded with its CRM decision before success is shown.')
req('MOD-C-008','Initial form scope','Launch forms collect inquiries and contact information. Public file uploads and payment fields require a later review of quotas, scanning, abuse and commercial behavior.')
req('MOD-C-009','Public pages and property content','Publish responsive, discoverable pages with search metadata. Property blocks must reflect provider capabilities and remove unavailable or disallowed listing content.')
h('Acceptance expectation')
p('An editor must be able to draft, preview, publish and roll back a page. A public form must create one correctly attributed inquiry. Domain ownership, keyboard editing, mobile layouts and listing attribution must be verified before release.')
h('Business boundary')
p('Public website content and the authenticated business workspace have different access needs. Drafts, customer records and private media must remain protected even when the public site shares the same brand.')

page('Property discovery and Trestle MLS','07 Property experience requirements')
p('Trestle is the selected MLS provider for the planned Las Vegas property experience. Approved board access, feed configuration and usage rights determine which live capabilities can be enabled. Preparation can proceed using visibly fictional property records.')
figure('mls-experience','Figure 4. Licensed property data connects to search and the agent workflow through an explicit policy boundary.')
req('MOD-F-001','Provider capabilities','Prepare search and inquiry workflows with a mock connection first. Distinguish native feed, live query and embedded experiences and disclose their feature limits.')
req('MOD-F-002','Property information','Normalize source identity, status, price, location, type, beds, baths, area, remarks, images, attribution and update time. Expose only permitted public information.')
req('MOD-F-003','Search and detail','Provide aligned map and list results, location and property filters, sorting, pagination, detail pages and inquiries. Hidden addresses must not leak precise location.')
req('MOD-F-004','Personal organizer','Let verified visitors save properties and searches within their site account. Alerts must respect availability, preferences and consent. Inquiries retain listing and source provenance.')
req('MOD-F-005','Freshness and withdrawal','Process updates reliably, prevent duplicate records and remove withdrawn or deleted listings across search, pages, images and scheduled alerts according to approved policy.')
req('MOD-F-006','Data usage policy','Control display, attribution, media, freshness, retention, caching, alerts and AI use independently. Unresolved permissions must disable the affected use.')
req('MOD-F-007','Live Trestle activation','Confirm the board, account, feed product, endpoint, available fields and participant rights at M6. Require mapping, partner acceptance and operating evidence before enabling live data.')
req('MOD-F-008','Clear demonstration data','Label fictional records and experiences as synthetic. Live production connections must reject mock feeds and synthetic records.')
h('Acceptance expectation')
p('Search must lead to a correctly routed inquiry. Live release also requires approved attribution, permitted images, measurable freshness and verified withdrawal across all display surfaces.')

page('Marketing and controlled automation','08 Engagement requirements')
p('The marketing module must support consistent follow up while giving users control over audiences, content, timing and delivery. A scheduled campaign remains subject to current consent and channel readiness when each action is due.')
figure('marketing-flow','Figure 5. Campaign execution includes an eligibility check at dispatch and confirmation of actual provider outcomes.')
req('MOD-E-001','Designs and campaigns','Provide reusable email and SMS designs, audience previews, broadcasts, scheduled sends, drip campaigns and birthday or transaction anniversary greetings. Licensed editor access must be verified.')
req('MOD-E-002','Workflow design','Support waits, messages, conditions and tasks in a bounded workflow without cycles. Preview timing and identify missing channel or subscription prerequisites before publication.')
req('MOD-E-003','Enrollment and versions','Support inquiry, saved search, favorite, stage and anniversary triggers. Define reentry policies and preserve the campaign version used by each enrollment.')
req('MOD-E-004','Reliable scheduled actions','Retain schedules through service interruptions. Recheck access, subscription, recipient status, consent, suppression, sender readiness and quiet hours before execution.')
req('MOD-E-005','Channel readiness and status','Show verified sender readiness and distinguish submission, delivery, bounce and failure. Development transports must deliver no real messages. Email opens are an imperfect engagement signal.')
req('MOD-E-006','Uncertain delivery','Prevent duplicate native actions and respect provider limits. Reconcile uncertain provider acceptance before retrying instead of assuming a message was never sent.')
req('MOD-E-007','Opt outs and cancellation','Record withdrawals and suppression promptly. Cancel eligible pending actions and disclose when a provider already accepted a message that cannot be recalled.')
req('MOD-E-008','Time and calendar rules','Apply recipient timezone or a visible fallback, quiet hours, daylight saving behavior and an explicit leap day policy for anniversary schedules.')
h('Acceptance expectation')
p('A consent withdrawal between audience preview and execution must block the send. Interrupted schedules must recover safely, and displayed delivery status must reflect provider evidence.')

page('Business reporting and attribution','09 Measurement requirements')
p('Leaders need consistent definitions and traceable results. Reporting must explain the selected scope, period, source and freshness so that performance discussions use comparable information.')
req('MOD-G-001','Defined business measures','Report website engagement, captured and qualified leads, contact rate, stage conversion, campaign outcomes, pipeline value and listing engagement. Separate counts, unique people, cohorts and estimates.')
req('MOD-G-002','Consistent definitions','Version metric definitions, including event source, calculation, exclusions, deduplication, timezone and freshness. Preserve the basis of historical reports.')
req('MOD-G-003','Transparent attribution','Use a documented conversion window and show unattributed activity. Email opens must not be represented as confirmed reading or business conversion.')
req('MOD-G-004','Scoped and privacy aware reporting','Apply the same access boundaries as the underlying customer records. Keep full contacts and raw inquiry content out of anonymous traffic measurement.')
req('MOD-G-005','Source and availability','Show original source periods, refresh time and unavailable states. Directory publication coverage must remain separate from measured search or AI answer visibility.')
req('MOD-G-006','Useful analysis and exports','Provide period comparisons, funnel and pipeline views, campaign or listing detail and account health. Authorize large exports and provide clear progress and completion states.')
h('Executive reporting views')
table(['View','Business question','Required context'],[
('Lead funnel','Where do inquiries progress or stop','Source, stage definitions, period and scope'),
('Agent activity','Are assigned opportunities receiving follow up','Assignment, task and verified activity history'),
('Campaign performance','What was sent, delivered and attributed','Channel, consent eligibility and conversion window'),
('Property engagement','Which permitted listings prompt interest','Listing source, availability and event definition'),
('Operational health','Which services need intervention','Last success, current state and responsible action')],[1.42,2.7,2.88])
h('Acceptance expectation')
p('Dashboard totals, detailed rows and exports must reconcile for the same scope and period. Unavailable information must remain visible as unavailable rather than being converted to a misleading zero or success score.')

page('AI assistance with human control','10 Intelligence requirements')
p('AI assistance is intended to reduce drafting and review effort. It must use permitted facts, explain its basis and leave publication, communication and commercial decisions with authorized people.')
figure('ai-review','Figure 6. AI generates proposals; human acceptance saves a draft and does not perform publication or sending.')
req('MOD-H-001','Useful assistance','Offer optional page, blog and campaign drafts, property descriptions, lead summaries, workflow suggestions and report explanations. Track requester, purpose, usage and expiry.')
req('MOD-H-002','Reliable generation','Use validated structured responses and an evaluated model configuration. Refusals, incomplete output and failures must leave existing business records unchanged.')
req('MOD-H-003','Permitted sources','Retrieve facts within organization and user access. Listing content requires explicit AI usage rights, and citations must point to actual accessible sources.')
req('MOD-H-004','Reviewed acceptance','Recheck permission and destination version before saving a reviewed draft. AI must not autonomously publish, send, change consent or modify billing.')
req('MOD-H-005','Follow up support','Summarize voluntarily supplied intent, timeline and engagement. Do not infer protected characteristics or decide housing or credit eligibility.')
req('MOD-H-006','Later public assistant','Treat a public conversational assistant as a later optional feature. It may use published content and permitted listings, with citations and a normal consent aware inquiry handoff.')
req('MOD-H-007','Grounded explanations','Explain computed metrics using their real definitions and values. Prevent unapproved retrieval, external actions and access to another organization through malicious content.')
req('MOD-H-008','Cost and data controls','Set organization usage caps, bounded input and output, retention settings and redacted monitoring. Approve provider data handling before activation.')
h('Acceptance expectation')
p('Before production activation, evaluations must demonstrate factual grounding, appropriate abstention, privacy, access boundaries, useful output, cost control and safe handling of malicious or incomplete input.')

page('Brand and local business presence','11 Presence requirements')
p('The platform must keep the customer brand consistent across websites and campaigns. Provider supported business presence features can extend this record to external channels when commercial rights and account permissions are confirmed.')
req('MOD-I-001','Shared business profile','Maintain business identity, contact details, hours, categories, media and social links. Verify network provider rights and supported functions before connecting directories.')
req('MOD-I-002','Controlled synchronization','Publish profile changes through versioned jobs with visible validation, provider errors, retries and last success. Keep listings, duplicates and reviews distinct.')
req('MOD-I-003','Honest presence measurement','Report directory coverage, reputation and source specific visibility using clear definitions. AI answer visibility requires a separate explicit measurement capability.')
req('MOD-I-004','External account workflows','Gate Google Business, social scheduling and review responses by provider support and user permission. AI review responses require human approval before publication.')
req('MOD-I-005','Reusable brand assets','Include brand kits and reusable page and email templates. Defer print fulfillment, standalone design packs, mortgage products and advanced creative editors pending demand and capability validation.')
h('Feature availability boundaries')
table(['Capability','Product treatment','Activation condition'],[
('Brand kits and reusable templates','Included in the defined launch product','Product acceptance and customer entitlement'),
('Directory profile synchronization','Provider dependent','Contract, supported API and customer authorization'),
('Business account or social connection','Optional and gated','Verified provider capability and permissions'),
('Review response publication','Human controlled','Approved content and supported publication method'),
('Print and advanced creative services','Deferred','Separate scope, licensing and fulfillment decision')],[2.1,2.1,2.8])
h('Acceptance expectation')
p('An external provider error must remain visible and actionable without silently changing the primary business profile. Coverage counts must describe publication status accurately and must not imply demonstrated market reach or AI recommendations.')

page('Support and service visibility','12 Service requirements')
p('Customers and support staff need a truthful view of account and integration health. The product must explain what happened, whether intervention is required and which authorized action can resolve the issue.')
req('MOD-J-001','Support tickets','Create scoped tickets with subject, category and description. Preserve the actor and status timeline. Add attachments only after the upload and retention capability is accepted.')
req('MOD-J-002','Customer safe service status','Show queued, running, succeeded, failed and unknown states with attempts, last success and a safe recommended action. Never expose secrets or another customer record.')
req('MOD-J-003','Operational support','Provide staff views for provisioning, billing, feeds, sender readiness, domains and job delays. Replays must be explicit, audited and safe against duplicate actions.')
req('MOD-J-004','Relevant help','Provide onboarding and contextual guidance that reflects the actual configured feature and provider status.')
h('Service state vocabulary')
table(['State','What the user should understand'],[
('Unconfigured','An account, setting or provider connection is still required'),
('Pending authorization','Permission or commercial approval is awaiting completion'),
('Queued or running','The requested operation has been accepted but is unfinished'),
('Ready or succeeded','The relevant setup or operation has completed with evidence'),
('Degraded or failed','An issue affects service; show a safe next action'),
('Unknown','The provider outcome is uncertain and must be reconciled')],[2.0,5.0])
h('Acceptance expectation')
p('A customer must be able to track a support request and understand an affected integration. Staff intervention must record the purpose and affected resource, enforce access expiry and preserve customer isolation.')
h('Operating responsibilities')
p('The launch plan must name owners for customer support, incident response, billing reconciliation, property feed freshness, delivery reconciliation and recovery. Contextual help must remain current as capabilities are activated or restricted.')

page('A complete application on iPhone and Android','13 Mobile application requirements')
p('Agent8lite must function as a full web application and an installable Progressive Web App on iPhone and Android. Customers can add it to their home screen and continue using the browser. App store distribution and a native wrapper are outside the defined architecture.')
figure('mobile-experience','Figure 7. Conceptual mobile workflow examples. Core functions require usable mobile paths, including accessible alternatives for complex editors.',width=5.65)
req('MOD-K-001','Installable application','Provide secure delivery, original icons, a stable application identity and standalone launch. Browser use must remain available without installation.')
req('MOD-K-002','Installation guidance','Offer appropriate iPhone and Android guidance based on available browser features and installed state. Do not block use behind installation.')
req('MOD-K-003','Complete core workflows','Support CRM, tasks, dashboards, website editing and publication, campaigns, property search, reports, settings and authorized administration at 360 pixels and above. Complex editors require usable alternatives.')
req('MOD-K-004','Device behavior','Respect safe areas, orientation, screen readers, text enlargement and virtual keyboards. Preserve navigation and safe return from sign in and hosted billing.')
req('MOD-K-005','Online first behavior','Cache approved static assets and a generic offline screen. Do not persist private customer or MLS data offline; clearly disable offline changes initially.')
req('MOD-K-006','Interrupted work','Warn about unsaved changes and refresh permissions and versions after reconnection. Durable offline editing is deferred; device process termination may discard unsaved tab memory.')
req('MOD-K-007','Safe updates','Notify users of application updates without discarding drafts or interrupting confirmation. Apply updates after acceptance or an appropriate idle state.')
req('MOD-K-008','Optional notifications','Offer generic opt in Web Push where supported. The in app inbox remains the dependable record when push is unavailable or denied.')
req('MOD-K-009','Device access control','Protect device subscriptions and revoke them on removal, logout or membership revocation. Notification links must recheck authentication and permissions.')
req('MOD-K-010','Physical device acceptance','Verify browser and installed modes on physical iPhone and Android devices, including all core workflows, interruptions, upgrades, push settings and private cache checks.')

page('Shared quality and trust requirements','14 Product standards')
p('These standards apply across the product. They are release requirements and proposed operating targets, not a statement of performance already achieved or a contractual service commitment.')
table(['Requirement','Required product behavior'],[
('Access and privacy','Separate organizations and own or team access. Revoke access across screens, jobs and downloads. Protect stored and transmitted data and keep sensitive details out of general logs.'),
('Clear actions and history','Use explicit confirmation for financial, publication and messaging actions. Preserve revisions and audit history. Repeated submissions must not duplicate business changes.'),
('Accessibility','Target the Web Content Accessibility Guidelines (WCAG) 2.2 AA with keyboard access, visible focus, readable errors, screen readers, accessible sign in and usable mobile layouts.'),
('Reliable recovery','Retain committed work and schedules through interruptions. Show pending and uncertain outcomes honestly. Validate recovery and rollback before rollout.'),
('Provider rights and budgets','Gate licensed uses and external actions by verified capability, permission and subscription. Apply quotas and spend limits so one organization cannot exhaust shared capacity.'),
('Data governance','Define retention, export, deletion and incident procedures. Keep customer data out of unapproved analytics and AI contexts. Preserve consent and source provenance.')],[1.5,5.5])
h('Proposed operating targets for approval')
table(['Measure','Planning target','Launch evidence'],[
('Availability','99.9 percent monthly','Approved measurement and dependency treatment'),
('Ordinary workspace reads','95th percentile at or below 400 ms','Representative load and warm read scenarios'),
('Property queries','95th percentile at or below 800 ms','Agreed filters and load; separate provider latency'),
('Mobile experience','Content load at or below 2.5 s; interaction response at or below 200 ms; layout shift score at or below 0.1','75th percentile on representative connectivity'),
('Ordinary capture notices','Within 60 seconds','Measured capture and notification workflow'),
('Recovery','Data loss window at or below 15 minutes; restoration within 4 hours','A demonstrated restore and integrity check')],[1.6,2.58,2.82])
p('Property refresh and withdrawal timing must follow the approved feed policy. Messaging timing must follow consent, quiet hours and channel policy. Executive owners must approve the test load, operating budget and final service commitments before release.')

page('Feature scope and release treatment','15 Capability portfolio')
p('The portfolio below describes the intended product and its delivery stages. The milestones are acceptance gates. No row implies that a feature has already passed runtime acceptance.')
table(['Capability','Core scope','Delivery stage','Release dependency'],[
('Accounts and billing','Organizations, roles, subscriptions and invoices','M1 to M2','Identity and payment configuration'),
('Agent workspace','Activity, tasks, appointments and notices','M2','Shared permissions and scheduling events'),
('CRM and routing','Contacts, leads, pipeline, imports and exports','M2','Verified ownership and data isolation'),
('Website builder','Templates, pages, domains, forms and publishing','M3','Content safety and domain verification'),
('Marketing','Email, SMS, campaigns and automation','M4','Sender readiness, consent and channel authorization'),
('Property experience','Search, detail, organizer and inquiries','M5 preparation; M6 live','Approved Trestle feed and permitted uses'),
('Reporting','Defined metrics, attribution and business views','M2 basics; M7 expansion','Source definitions and reconciliation'),
('AI assistance','Reviewed drafts, summaries and explanations','M7','Evaluation, provider agreement and data rights'),
('Brand and presence','Brand kits, profiles and supported sync','M3 templates; M7 sync','External capabilities and permissions'),
('Support and operations','Tickets, health, help and controlled recovery','M2 basics; M8 readiness','Operational ownership and audit'),
('Mobile application','Full browser and installed iPhone or Android paths','M1 design; throughout delivery; M8 acceptance','Physical device and update verification')],[1.4,2.1,1.4,2.1],compact=True)
h('Deferred capabilities')
p('Durable offline editing, a public conversational assistant, public form uploads, public payment fields, support attachments, print fulfillment, mortgage products and advanced standalone creative tools require later scope and capability decisions. Social and external review publication remain provider dependent.')
h('Commercial packaging')
p('Plan tiers, included usage, premium services and overage treatment must be approved separately. The product must show actual entitlements and pending provider activation rather than presenting all portfolio capabilities as immediately available to every customer.')

page('Delivery roadmap and prerequisites','16 Delivery plan')
p('The eight milestones organize delivery around complete workflows and evidence. Teams may work in parallel where interfaces and prerequisites are settled. Dates and investment levels require a resourced delivery plan.')
figure('delivery-roadmap','Figure 8. Delivery milestones in their principal acceptance order. This is a sequence, not a calendar commitment.')
h('Dependencies that shape delivery')
bullet('M1 establishes secure identity, organization boundaries and shared standards. M2 enables customer records and the accountable agent workflow.')
bullet('M3 connects public content and forms to CRM. M4 adds controlled engagement. M5 proves the property experience with fictional data and can overlap eligible marketing work.')
bullet('M6 activates live Trestle only after licensed access, mapping, policy and partner acceptance are complete.')
bullet('M7 expands reporting, assistance and presence. Non MLS work can proceed while live feed approval is pending.')
bullet('M8 verifies launch scope, real devices, accessibility, performance, restore, support ownership and controlled release.')
p('A clearly defined CRM and website release can be considered without live IDX if its scope is explicit and its own acceptance conditions pass. A release marketed with live MLS search requires M6 acceptance.')

page('Launch acceptance and executive release decision','17 Acceptance requirements')
p('Release approval must be based on demonstrated customer journeys and operating evidence. Each scenario below must pass for the capabilities included in the launch scope.')
table(['Scenario','Evidence required for acceptance'],[
('Onboard and invite','Correct organization access, verified invitation, last owner protection and accurate capability status'),
('Capture and follow up','One inquiry decision, explainable eligible assignment, complete source history and an actionable follow up task'),
('Draft and publish a website','Protected drafts, working preview, reviewed publication, safe rollback, verified domains and accessible mobile pages'),
('Find a property and inquire','Consistent search and detail, permitted attribution and images, private organizer and correctly routed CRM inquiry'),
('Schedule a campaign','Reviewed audience and message, current consent at dispatch, quiet hours, safe recovery and provider based status'),
('Review AI assistance','Grounded sources, proper permissions, evaluated output and human controlled draft acceptance'),
('Reconcile a report','Matching scoped totals and detail, definitions, timezones, refresh status and traceable export'),
('Work on physical phones','Browser and home screen use on iPhone and Android; core actions, interruptions, safe updates and private cache checks'),
('Recover and support','A verified restore, incident and rollback procedure, service monitoring, access expiry and named operational owners')],[2.0,5.0])
h('Release approval conditions')
bullet('All included functional requirements and shared access standards have acceptance evidence; no critical organization isolation or data rights defect remains.')
bullet('Live providers have the appropriate credentials, commercial rights, customer permissions and operating limits for their enabled functions.')
bullet('Performance and recovery targets are measured against the approved load and operating environment.')
bullet('Customer communication, support ownership, phased rollout and rollback decisions are approved by the release owner.')
p('Document completeness, interface concepts and generated specifications do not substitute for application testing. The release record must identify the exact enabled scope and any deferred capability.')

page('Business measures and executive decisions','18 Business governance')
p('Business results should be measured against an agreed baseline after instrumentation is validated. Targets belong to the commercial operating plan and must not be inferred from illustrative interfaces or implementation progress.')
table(['Measure','Definition to approve','Business use'],[
('Activation','Organizations completing the agreed onboarding and first useful workflow','Identify setup friction'),
('Lead response','Elapsed time from captured inquiry to the first verified qualifying agent action','Assess follow up execution'),
('Lead progression','Unique leads reaching a defined next stage within an agreed cohort and period','Assess pipeline continuity'),
('Workflow adoption','Active eligible users completing selected meaningful product actions','Assess product utility'),
('Campaign contribution','Eligible conversions under the defined attribution window, including unattributed activity','Assess engagement outcomes'),
('Retention and unit economics','Commercially agreed customer retention, recurring revenue and cost definitions','Assess commercial sustainability'),
('Service health','Availability, incident handling, provider freshness and demonstrated recovery','Assess operating reliability')],[1.55,3.3,2.15],compact=True)
h('Executive decisions needed')
table(['Decision','Accountable function','Required before'],[
('Launch audience, region and included feature scope','Product and commercial leadership','Resourced delivery commitment'),
('Plan tiers, pricing, usage limits and service commitments','Commercial and finance leadership','Customer subscription launch'),
('Trestle account, board access and permitted uses','Business owner with MLS or provider partners','M6 live activation'),
('Messaging channels, sender accounts and approved policy','Marketing and operations with appropriate advisers','Real outbound messaging'),
('AI provider terms, retention and acceptable use','Product, privacy and security leadership','M7 production AI activation'),
('Hosting region, capacity, support model and budget','Technology and operations leadership','Production sizing and M8 approval'),
('Measurement baseline, targets and attribution definitions','Commercial and analytics owners','Business performance reporting')],[3.2,2.2,1.6],compact=True)
p('Named individuals, budgets and deadlines must be assigned in the delivery plan. These functions describe responsibilities and do not imply that appointments or approvals have already been made.')

page('Requirement coverage and document governance','19 Reference and control')
p('This executive specification translates the approved technical blueprint into customer and business requirements. The requirement identifiers preserve traceability to the implementation specification. Detailed APIs, data structures and engineering controls remain in the Technical Design Document.')
table(['Feature family','Source identifiers','Count','Document page'],[
('Accounts and billing','MOD A 001 through 006',6,5),
('Agent workspace','MOD B 001 through 005',5,5),
('CRM and routing','MOD D 001 through 007',7,6),
('Websites and forms','MOD C 001 through 009',9,7),
('Property discovery and MLS','MOD F 001 through 008',8,8),
('Marketing and automation','MOD E 001 through 008',8,9),
('Reporting and attribution','MOD G 001 through 006',6,10),
('AI assistance','MOD H 001 through 008',8,11),
('Brand and local presence','MOD I 001 through 005',5,12),
('Support and operations','MOD J 001 through 004',4,13),
('Mobile application','MOD K 001 through 010',10,14),
('Total functional coverage','11 feature families',76,'5 to 14')],[2.05,2.6,.65,1.7])
h('Baseline and governing records')
bullet('Competitive Real Estate SaaS Technical Design Document version 1.0, dated 10 October 2026, and its structured technical design record.')
bullet('Owner requirements: Trestle as the MLS provider and a complete browser and home screen application for iPhone and Android without app store distribution.')
bullet('Implementation status dated 10 October 2026: development active; no milestone recorded as complete; live MLS, messaging, payments and production rollout not enabled.')
h('Change and approval process')
p('Changes to requirements, scope, operating targets or provider assumptions must be recorded with the business reason, affected capabilities and acceptance impact. Product, technology and the relevant commercial or operational owner review the change before it becomes a release commitment.')
p('The technical design governs engineering details. Resolve ambiguity through a recorded product decision against the approved specification. Provider capability and usage rights continue to govern activation.')

expected=[r['id'] for m in TDD['modules'] for r in m['requirements']]
assert len(page_titles)==20,(len(page_titles),page_titles)
assert len(req_ids)==76 and set(req_ids)==set(expected),{'missing':set(expected)-set(req_ids),'extra':set(req_ids)-set(expected)}
out=ROOT/'Agent8lite-Application-Requirements-and-Features-2026-10-10.docx'
for element in [doc._element,doc.styles._element]:
    for border in element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.save(out)
manifest={'title':doc.core_properties.title,'date':'2026-10-10','pagesPlanned':20,'pageTitles':page_titles,'requirementCount':len(req_ids),'requirementIds':req_ids,'featureFamilies':11,'visuals':9,'source':'blueprint/technical-design.json','statusBaseline':'IMPLEMENTATION_STATUS.md','deliverables':[out.name,out.with_suffix('.pdf').name]}
(ROOT/'document-manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print(json.dumps({'docx':str(out),'pagesPlanned':20,'requirements':76,'visuals':9}))
